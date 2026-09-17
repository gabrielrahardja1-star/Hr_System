"""Talenta 'Import Attendance' skeleton enrichment (Phase 5)."""

from __future__ import annotations

import datetime as dt

import pytest

WIB = dt.timezone(dt.timedelta(hours=7))

TALENTA_HEADERS = [
    "Employee ID*", "Full Name", "Date*", "Shift", "Shift Code", "Shift Label",
    "Schedule In", "Schedule Out", "Attendance Code", "Check In", "Check Out",
    "Overtime Check In", "Overtime Check Out", "Overtime Before", "Overtime After",
    "Overtime Break Before", "Overtime Break After", "Time Off Code",
    "Time Off Halfday", "Time Off Schedule In", "Time Off Schedule Out",
    "Branch", "Organization", "Job Position", "Job Level", "Employment Status",
    "Join Date", "Lock Status",
]


def _write_skeleton(path, rows: list[dict]):
    """rows: {talenta_id, name, date, shift} — 'S2' working or 'dayoff'/'National Holiday'."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Export Attendance"
    ws.append(TALENTA_HEADERS)
    for r in rows:
        working = r["shift"] == "S2"
        line = [""] * len(TALENTA_HEADERS)
        line[0] = r["talenta_id"]
        line[1] = r["name"]
        line[2] = r["date"]
        line[3] = r["shift"]
        line[6] = "07:20" if working else "00:00"
        line[7] = "16:00" if working else "00:00"
        line[17] = r.get("time_off_code", "")
        line[21] = "PT Cinta Kerja Indonesia (A)"
        line[27] = 0
        ws.append(line)
    wb.save(path)
    return path


@pytest.fixture()
def skeleton_factory(tmp_path):
    def _make(rows, name="skeleton.xlsx"):
        return _write_skeleton(tmp_path / name, rows)
    return _make


def _emp(make_employee, talenta_id, device_user_id="9001"):
    return make_employee(
        device_user_id=device_user_id,
        emp_code=f"E-{device_user_id}",
        talenta_id=talenta_id,
        name="Test Worker",
        department="Operations",
        shift_key="KERJA",
        roster_pattern="continuous",
    )


def _period_days(period="2026-08"):
    from server.core.period import talenta_period_bounds

    start, end = talenta_period_bounds(period)
    return start, end


def test_period_bounds_are_26_to_25():
    from server.core.period import talenta_period_bounds, talenta_period_for

    assert talenta_period_bounds("2026-08") == (dt.date(2026, 7, 26), dt.date(2026, 8, 25))
    assert talenta_period_bounds("2026-01") == (dt.date(2025, 12, 26), dt.date(2026, 1, 25))
    assert talenta_period_for(dt.date(2026, 7, 26)) == "2026-08"
    assert talenta_period_for(dt.date(2026, 7, 25)) == "2026-07"


def test_passthrough_non_working_rows(make_employee, add_punches, session, skeleton_factory):
    from server.core.talenta_export import enrich_skeleton

    _emp(make_employee, "CKI-A230065")
    sk = skeleton_factory([
        {"talenta_id": "CKI-A230065", "name": "T", "date": "2026-07-27", "shift": "dayoff"},
        {"talenta_id": "CKI-A230065", "name": "T", "date": "2026-08-17", "shift": "National Holiday"},
    ])
    r = enrich_skeleton(session, sk, "2026-08", kind="final")
    assert r.passthrough == 2
    assert r.enriched == 0 and r.held == 0
    assert not r.blocked


def test_enriches_a_clean_working_day(make_employee, add_punches, session, skeleton_factory):
    from server.core.recompute import recompute_all
    from server.core.talenta_export import enrich_skeleton
    from openpyxl import load_workbook

    emp = _emp(make_employee, "CKI-A230065")
    add_punches(emp.device_user_id, [
        dt.datetime(2026, 7, 28, 8, 18, tzinfo=WIB),
        dt.datetime(2026, 7, 28, 16, 5, tzinfo=WIB),
    ])
    recompute_all(session, *_period_days(), respect_frontier=False)
    session.commit()

    sk = skeleton_factory([
        {"talenta_id": "CKI-A230065", "name": "T", "date": "2026-07-28", "shift": "S2"},
    ])
    r = enrich_skeleton(session, sk, "2026-08", kind="final")
    assert r.enriched == 1 and not r.blocked

    ws = load_workbook(r.out_path).active
    hdr = {c.value: c.column for c in ws[1]}
    row2 = 2
    assert ws.cell(row=row2, column=hdr["Attendance Code"]).value == "H"
    assert ws.cell(row=row2, column=hdr["Check In"]).value == "08:18"
    assert ws.cell(row=row2, column=hdr["Check Out"]).value == "16:05"
    # untouched columns stay put
    assert ws.cell(row=row2, column=hdr["Schedule In"]).value == "07:20"
    assert ws.cell(row=row2, column=hdr["Branch"]).value == "PT Cinta Kerja Indonesia (A)"


def test_missing_punch_day_is_held_and_blocks_final(make_employee, add_punches, session, skeleton_factory):
    from server.core.recompute import recompute_all
    from server.core.talenta_export import enrich_skeleton

    emp = _emp(make_employee, "CKI-A230065")
    add_punches(emp.device_user_id, [dt.datetime(2026, 7, 28, 8, 20, tzinfo=WIB)])  # in only
    recompute_all(session, *_period_days(), respect_frontier=False)
    session.commit()

    sk = skeleton_factory([
        {"talenta_id": "CKI-A230065", "name": "T", "date": "2026-07-28", "shift": "S2"},
    ])
    final = enrich_skeleton(session, sk, "2026-08", kind="final")
    assert final.blocked and final.held == 1 and final.out_path is None
    assert "MP" in final.held_rows[0].reason

    draft = enrich_skeleton(session, sk, "2026-08", kind="draft")
    assert not draft.blocked and draft.out_path is not None and draft.held == 1
    from openpyxl import load_workbook
    ws = load_workbook(draft.out_path).active
    hdr = {c.value: c.column for c in ws[1]}
    assert ws.cell(row=2, column=hdr["Check In"]).value in (None, "")


def test_unmapped_employee_is_held(make_employee, session, skeleton_factory):
    from server.core.talenta_export import enrich_skeleton

    # employee exists in the skeleton but no talenta_id match in our DB
    sk = skeleton_factory([
        {"talenta_id": "CKI-A999999", "name": "Ghost", "date": "2026-07-28", "shift": "S2"},
    ])
    r = enrich_skeleton(session, sk, "2026-08", kind="draft")
    assert r.held == 1
    assert "not mapped" in r.held_rows[0].reason


def test_time_off_rows_pass_through(make_employee, session, skeleton_factory):
    from server.core.talenta_export import enrich_skeleton

    _emp(make_employee, "CKI-A230065")
    sk = skeleton_factory([
        {"talenta_id": "CKI-A230065", "name": "T", "date": "2026-07-28",
         "shift": "S2", "time_off_code": "CT"},
    ])
    r = enrich_skeleton(session, sk, "2026-08", kind="final")
    assert r.passthrough == 1 and r.held == 0 and not r.blocked
