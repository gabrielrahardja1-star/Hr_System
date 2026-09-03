"""Enrich a Talenta 'Import Attendance' skeleton with reviewed punch data.

Input : a .xlsx exported FROM Talenta for one 26→25 payroll period — one row
        per employee per day, with Shift / Schedule / org columns already filled
        and Attendance Code / Check In / Check Out blank.
Output: the same workbook (styles + comments preserved) with, for every working
        day we can finalise:
            Attendance Code  <- code_map[our day code]   (e.g. "H")
            Check In         <- first punch, HH:MM site-local
            Check Out        <- last punch,  HH:MM site-local

We never touch a non-working row (dayoff / National Holiday / time-off) and we
never touch any other column. A working row we cannot finalise (missing punch,
absent-pending-code, employee not mapped, no punch data) is HELD: a `final` run
refuses to write while any row is held; a `draft` run writes the rest and lists
them.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from server.config import CONFIG_DIR, _load_yaml, get_settings, get_wage_mapping
from server.core.period import talenta_period_bounds
from server.models import DayRecord, Employee, Exception_, ExceptionState


def _cfg(path: str | None = None) -> dict:
    return _load_yaml(Path(path) if path else CONFIG_DIR / "talenta_export.yaml")


@dataclass
class RowOutcome:
    row: int
    employee_id: str
    date: str
    status: str              # "enriched" | "held" | "passthrough" | "skipped"
    code: str | None = None
    check_in: str | None = None
    check_out: str | None = None
    reason: str | None = None


@dataclass
class EnrichResult:
    period: str
    kind: str                # "draft" | "final"
    skeleton_rows: int
    enriched: int = 0
    held: int = 0
    passthrough: int = 0
    skipped: int = 0
    out_path: Path | None = None
    blocked: bool = False
    outcomes: list[RowOutcome] = field(default_factory=list)

    @property
    def held_rows(self) -> list[RowOutcome]:
        return [o for o in self.outcomes if o.status == "held"]


def _header_map(ws, columns_cfg: dict) -> dict[str, int]:
    """role -> 1-based column index, matched case-insensitively on row 1."""
    header = {
        str(c.value).strip().lower(): c.column
        for c in ws[1]
        if c.value is not None
    }
    out: dict[str, int] = {}
    for role, text in columns_cfg.items():
        idx = header.get(str(text).strip().lower())
        if idx is None and role in ("employee_id", "date", "shift",
                                    "attendance_code", "check_in", "check_out"):
            raise ValueError(
                f"Skeleton is missing required column {text!r} (role {role})"
            )
        if idx is not None:
            out[role] = idx
    return out


def _as_date_str(value) -> str:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()[:10]


def enrich_skeleton(
    session: Session,
    skeleton_path: str | Path,
    period: str,
    *,
    kind: str = "draft",
    out_path: str | Path | None = None,
    config_path: str | None = None,
) -> EnrichResult:
    cfg = _cfg(config_path)
    cols = cfg["columns"]
    non_working = {str(v).strip().lower() for v in cfg.get("non_working_shift_values", [])}
    code_map: dict[str, str] = dict(cfg.get("code_map", {}))
    default_code = cfg.get("default_present_code", "H")
    time_fmt = cfg.get("time_format", "%H:%M")
    hold_when_no_record = bool(cfg.get("hold_when_no_record", True))
    tz = get_settings().timezone

    start, end = talenta_period_bounds(period)

    skeleton_path = Path(skeleton_path)
    wb = load_workbook(skeleton_path)
    ws = wb.active
    h = _header_map(ws, cols)

    # Preload our data for the period.
    emp_by_talenta = {
        e.talenta_id: e
        for e in session.execute(
            select(Employee).where(Employee.talenta_id.is_not(None))
        ).scalars()
    }
    recs = session.execute(
        select(DayRecord).where(
            DayRecord.work_date >= start, DayRecord.work_date <= end
        )
    ).scalars().all()
    rec_by_key = {(r.employee_id, r.work_date.isoformat()): r for r in recs}

    wage = get_wage_mapping()
    blocking_exc: dict[int, str] = {}
    for drid, kind_ in session.execute(
        select(Exception_.day_record_id, Exception_.kind).where(
            Exception_.state == ExceptionState.open,
            Exception_.day_record_id.in_([r.id for r in recs] or [0]),
        )
    ):
        k = kind_.value
        if wage.code_blocks_export(k) or wage.flag_blocks_export(k):
            blocking_exc[drid] = k

    result = EnrichResult(period=period, kind=kind, skeleton_rows=ws.max_row - 1)

    for row in range(2, ws.max_row + 1):
        eid = ws.cell(row=row, column=h["employee_id"]).value
        if eid is None or str(eid).strip() == "":
            result.skipped += 1
            result.outcomes.append(RowOutcome(row, "", "", "skipped", reason="blank row"))
            continue
        eid = str(eid).strip()
        date_str = _as_date_str(ws.cell(row=row, column=h["date"]).value)
        shift_val = str(ws.cell(row=row, column=h["shift"]).value or "").strip()

        # Non-working day -> leave untouched.
        if shift_val.lower() in non_working:
            result.passthrough += 1
            result.outcomes.append(RowOutcome(row, eid, date_str, "passthrough",
                                              reason=shift_val))
            continue

        # Row already carries a Time Off code -> Talenta owns leave, leave it.
        if "time_off_code" in h and ws.cell(row=row, column=h["time_off_code"]).value:
            result.passthrough += 1
            result.outcomes.append(RowOutcome(row, eid, date_str, "passthrough",
                                              reason="time off"))
            continue

        emp = emp_by_talenta.get(eid)
        rec = rec_by_key.get((emp.id, date_str)) if emp else None

        def hold(reason: str) -> None:
            result.held += 1
            result.outcomes.append(RowOutcome(row, eid, date_str, "held", reason=reason))

        if emp is None:
            hold("employee not mapped to a device user (set talenta_id)")
            continue
        if rec is None:
            if hold_when_no_record:
                hold("no reviewed day record — punches not synced / not computed")
            else:
                result.skipped += 1
                result.outcomes.append(RowOutcome(row, eid, date_str, "skipped",
                                                  reason="no record"))
            continue

        # The skeleton's Shift column already told us this is a working day, so
        # the only question is whether we can finalise it:
        #   open blocking exception (MP/SS/...)  -> hold (same gate as the UI)
        #   no in/out pair (absent, single punch) -> hold
        #   otherwise                             -> H (or a code_map override)
        if rec.id in blocking_exc:
            hold(f"open {blocking_exc[rec.id]} exception — resolve it first")
            continue
        if rec.first_in is None or rec.last_out is None:
            if rec.code == "A":
                hold("no punches on a Talenta working day — needs the absent code")
            elif rec.code in ("WO", "H"):
                hold(
                    "Talenta scheduled this day but our roster had it off, and no "
                    "punches were recorded — confirm the schedule / mark absent"
                )
            else:
                hold(f"no in/out pair (day is {rec.code})")
            continue
        mapped = code_map.get(rec.code, default_code)

        ci = rec.first_in.astimezone(tz).strftime(time_fmt)
        co = rec.last_out.astimezone(tz).strftime(time_fmt)
        result.outcomes.append(
            RowOutcome(row, eid, date_str, "enriched", code=mapped,
                       check_in=ci, check_out=co)
        )
        result.enriched += 1

    # A final run must not write a partial file.
    if kind == "final" and result.held:
        result.blocked = True
        return result

    # Apply the enriched cells.
    for o in result.outcomes:
        if o.status != "enriched":
            continue
        ws.cell(row=o.row, column=h["attendance_code"]).value = o.code
        ws.cell(row=o.row, column=h["check_in"]).value = o.check_in
        ws.cell(row=o.row, column=h["check_out"]).value = o.check_out

    if out_path is None:
        settings = get_settings()
        settings.export_dir.mkdir(parents=True, exist_ok=True)
        suffix = "final" if kind == "final" else "draft"
        out_path = settings.export_dir / f"talenta_import_{period}_{suffix}.xlsx"
    out_path = Path(out_path)
    wb.save(out_path)
    result.out_path = out_path
    return result
